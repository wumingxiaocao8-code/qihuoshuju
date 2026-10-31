"""开盘复盘：每次开盘后把每个组合开盘那一针记一行，对答案，出记分卡。

平时不用手动跑：开盘tick.py 取完 tick（21:03、9:03）会调用这里，把「数据/」里所有开盘重新算一遍。
手动跑（在本文件夹里）：python 开盘复盘.py

每次开盘、每个组合一行（复盘数据/开盘复盘.csv）：
  昨收价差    两腿上一次收盘价之差（开盘文件夹里的「收盘价.csv」，没有就用「预估.csv」里记的）
  针          竞价那张快照（20:59 / 8:59，成交手数 > 0）两腿的价差；有一腿竞价没成交就空着
  第一次真实成交  开盘后两腿第一次在同一个半秒里都有成交，那半秒的均价价差
  开盘后 1、5、30、60 秒的价差（最新价）；第一分钟里买价差最低、卖价差最高（买价差 = 腿1卖一 − 腿2买一）
  外盘        国内收盘 → 开盘前 7 分钟（20:53 / 8:53，跟预估同一时刻）的涨跌；外盘方向 涨 / 跌，和小草定的分档：
              0–1% 没波动，1–3% 小，3–6% 大，≥6% 超大
  同向 / 反向（小草定的）  外盘涨、价差（腿1 − 腿2，次主力 − 主力）比昨收价差往上 = 同向；外盘跌、价差往下也是同向；
              反过来是反向。针和第一次真实成交各判断一次；没有外盘的（不锈钢、外盘没取到）不判断。
  跳空        腿1跳空、主力跳空（腿2）= 竞价价 − 昨收，单位和价格一样；外盘涨跌% 还是百分比
  预估        腿1预估、腿2预估（开盘价）、预测区间（这次开盘价差 = 针的八成落在）、预测方向（价差比昨收往上 / 往下 / 不变），
              在范围内 = 针落没落在预测区间里；预估误差 = 腿2 竞价价 − 腿2预估（主力价格）
              20:53 / 8:53 开盘预估.py 估完就写进来（那时还没 tick，这几列和昨收、外盘先有，开盘后整行补全）
  预估误差    腿2 竞价价 − 开盘预估.py 估的价（没有预估就空着）
  59 分判断   竞价判断.py 在 20:59 / 8:59 判的：真（✅ 大概率整根开到针附近）/ 半真（▲ 开到一部分）/ 假（❌）/ 没针 / 不判断，
              判断得分（竞价量分 + 盘口分）、托住几成、猜首笔、首笔半（昨收价差 + 0.5 ×（猜首笔 − 昨收价差），ag、sn、ni 是 1/3；没有猜首笔就空着）、预计首笔区间、首笔在区间里；
              判断来源 = 实盘（当时写的 竞价判断.csv）或 回放（那次没跑竞价判断，用存下来的 tick 补算）；
              v1判断 = 同一根针 v1 会判成什么（sc 10-24 起换成混合版，对照用）；立刻能卖、立刻能买 = 竞价后盘口的价差（拟合 sc 首笔区间用）
  模拟盈亏    照浮层标签做一手（见 SIM_*）：重仓标记、模拟成交、模拟挂价、模拟成交价差、模拟成交秒、模拟出场价差、模拟出场、
              模拟档数、模拟每手毛元、模拟手续费、模拟每手净元；判 ❌ 的也假设照做（模拟成交写「❌假设…」），给记分卡数错过的
  标签        星期、周一早盘、节后第一次开盘、换月期（每月 8–15 号）、腿1 离到期几天；
              仓储提示、长期事件（预估时记下的，照抄 output/ 里期货速览的结果）
同样内容另存一份 复盘数据/开盘复盘.xlsx（列宽调好，给人用 Excel 看；程序读 .csv）。
系数历史（复盘数据/系数历史.csv）：每次开盘、每个品种一行（主力 = 腿2），开盘预估.py 用它算系数。
  开头的「回测」行是 2024-05 起两年多的夜盘开盘（日线 + 雅虎小时线，外盘算到 21:00），后面是每次实盘记的。
网页：复盘网页/开盘日期_时间.html（最近一次开盘、记分卡、涨跌幅统计、预估准不准、明细），每次开盘一份。
"""
import html
import math
import re
import site
from datetime import datetime, timedelta
from pathlib import Path

# 单独运行时也能找到日记库 venv 里装好的包（venv 的 python.exe 被智能应用控制拦了，用原版 Python 跑）
VENV_SITE = Path.home() / ".venvs" / "vault" / "Lib" / "site-packages"
if VENV_SITE.exists():
    site.addsitedir(str(VENV_SITE))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

import 外盘涨跌 as wp  # noqa: E402

HERE = Path(__file__).resolve().parent
DATA = HERE / "数据"
REVIEW = HERE / "复盘数据"
COEF_FILE = REVIEW / "系数历史.csv"
BASE_FILE = REVIEW / "针_日线基线.csv"
REC_FILE = REVIEW / "开盘复盘.csv"
PAGE_DIR = HERE / "复盘网页"  # 每次开盘后的复盘网页：复盘网页/2026-10-23_0900.html（按最近一次开盘命名）
FORECAST_DIR = HERE / "预估"  # 每次开盘的预估网页：预估/2026-10-22_2100.html
LEAD = timedelta(minutes=7)  # 预估在开盘前 7 分钟（20:53 / 8:53）
BINS = [(1, "没波动"), (3, "小"), (6, "大"), (math.inf, "超大")]  # 外盘涨跌幅绝对值的上限（%）
BUCKETS = [b for _, b in BINS]
DIRS = ("涨", "跌")  # 外盘方向
COEF_N, BAND_N, ALARM_N = 250, 60, 20  # 系数用最近 250 次，误差范围看最近 60 次，报警看最近 20 次
DAY_MIN = 20  # 日盘攒够 20 次才单独算系数，之前借夜盘的
WEEK = "一二三四五六日"
REC_COLS = ["开盘", "场次", "星期", "组合", "品种", "腿1", "腿2", "腿1相对主力", "腿1昨收", "腿2昨收", "昨收价差",
            "腿1竞价", "腿2竞价", "腿1竞价手数", "腿2竞价手数", "针", "针变化", "腿1跳空", "主力跳空", "针方向",
            "针真假", "针偏几档", "首笔偏几档", "腿1竞价占持仓‰", "腿1竞价持仓变化",
            "59分判断", "判断得分", "托住几成", "猜首笔", "首笔半", "预计首笔区间", "首笔在区间里", "判断来源",
            "首笔秒", "首笔价差", "首笔变化", "首笔方向", "价差1秒", "价差5秒", "价差30秒", "价差60秒",
            "首分钟买价差最低", "首分钟卖价差最高", "外盘涨跌%", "外盘方向", "分档", "外盘时段", "外盘来源",
            "腿1预估", "腿2预估", "预测区间", "预测方向", "预估误差", "在范围内", "仓储提示", "长期事件",
            "周一早盘", "节后", "换月期", "腿1离到期天数", "备注", "v1判断", "立刻能卖", "立刻能买",
            "重仓标记", "模拟成交", "模拟挂价", "模拟成交价差", "模拟成交秒", "模拟出场价差", "模拟出场", "模拟档数",
            "模拟每手毛元", "模拟手续费", "模拟每手净元"]
PCT_COLS = ["主力跳空%", "腿1跳空%"]  # 只在程序里用（记分卡、涨跌幅统计），不写进表格
HIDDEN_COLS = PCT_COLS + ["预估下沿", "预估上沿", "主力在范围内", "首笔下沿", "首笔上沿"]  # 只在程序里用
# 写进表格的列和顺序（小草定的）：针变化只在程序里用（针方向、记分卡），不写进表格
OUT_COLS = ["开盘", "场次", "星期", "组合", "品种", "腿1", "腿2",
            "昨收价差", "针", "外盘涨跌%",
            "首笔价差", "首笔变化", "首笔方向", "价差1秒", "价差5秒", "价差30秒", "价差60秒", "首分钟买价差最低", "首分钟卖价差最高",
            "腿1竞价", "腿2竞价", "腿1竞价手数", "腿2竞价手数", "腿1跳空", "主力跳空",
            "针真假", "针偏几档", "首笔偏几档", "腿1竞价占持仓‰", "腿1竞价持仓变化",
            "59分判断", "判断得分", "托住几成", "猜首笔", "首笔半", "预计首笔区间", "首笔在区间里", "判断来源",
            "重仓标记", "模拟成交", "模拟成交价差", "模拟出场价差", "模拟档数", "模拟每手净元"]
OUT_COLS += [c for c in REC_COLS if c not in OUT_COLS and c != "针变化"]
# 针真不真（小草定的，2026-10-22）：首笔真实成交相对昨收价差偏出这么多档（最小变动价位）才算有动静
NEEDLE_GATE = {"sc": 15, "fu": 20, "bu": 20, "ag": 10, "cu": 6, "sp": 5, "nr": 5, "br": 5, "ru": 5, "ss": 5,
               "ao": 8, "hc": 5, "rb": 5, "al": 5, "pb": 5, "zn": 5,
               "sn": 25, "ni": 10}  # ao 10-31 从 5 改 8（ao 晃得大，5–7 档小针扣费后每手只剩约 1 档，见 分析/2026-10-31/其余品种回放/小结.md 第 5 节）；sn、ni 10-31 加：上期所两年回放，5 档时锡 86% 的开盘都算出门槛（盘口本身就宽），见 分析/2026-10-31/有色开盘/小结.md 第 6 节
# 收盘价.csv 没记最小变动价位时用这个
TICK = {"sc": 0.1, "fu": 1, "bu": 1, "ag": 1, "au": 0.02, "cu": 10, "sp": 2, "nr": 5, "br": 5, "ru": 5, "ss": 5,
        "ao": 1, "hc": 1, "rb": 1, "al": 5, "pb": 5, "zn": 5, "ni": 10, "sn": 10, "lu": 1}
# 腿1竞价占持仓的分档（‰）：10-22 从 3 次开盘看出来的，只是临时的线，攒多了再调
RATIO_BINS = [(0.7, "0.7‰ 以下"), (1.5, "0.7–1.5‰"), (math.inf, "1.5‰ 以上")]
TRUTHS = ("真", "半真", "假", "反向", "没针")
# 照浮层标签做的模拟盈亏（10-31 加，口径照抄两年回放 分析/2026-10-27/fu和bu回测/脚本/回测.py 和 测试/2026-10-31_首笔半效果/c.py）：
# 判 ✅▲ 的在首笔半挂一手，开盘后 SIM_VALID 秒内那半秒的成交均价差或立刻能成交的价比挂价好就成交（按那个更好的价算，不按挂价）；
# 止盈 = 挂价往有利方向 SIM_TP 档、止损 = 比成交价亏 SIM_SL 档，成交后 SIM_DELAY 秒内不止损，都按当时立刻能平的价平；
# tick 只到开盘后 60 秒，到时按最后能平的价算（两年回放最多拿 300 秒）。全市场的量都算成自己的，是上限。
SIM_VALID, SIM_TP, SIM_SL, SIM_DELAY = 5, 5, 10, 10
# 合约乘数（同 成交复盘.py 的 MULT），档数 × 最小变动价位 × 乘数 = 每手元
MULT = {"au": 1000, "ag": 15, "cu": 5, "bc": 5, "al": 5, "zn": 5, "pb": 5, "ni": 1, "sn": 1, "ao": 20,
        "rb": 10, "hc": 10, "ss": 5, "wr": 10, "ru": 10, "nr": 10, "br": 5, "bu": 10, "fu": 10, "lu": 10,
        "sc": 1000, "sp": 10}
# 每手每腿手续费（元）：开仓、平今。一组跨月开平 = 两腿 × (开 + 平今)。只填了从他成交记录算出来的；没填的品种只算毛的
# cu：开 27、平今 54 元/手 → 一组 162 元 ≈ 3.2 档；ao br sp rb hc ru ss 10-31 从他 10-22~10-28 的成交记录算（每手中位，
# 分析/2026-10-31/其余品种回放/脚本/分析.py 的 FEES；hc 按主力月，hc2703 那腿其实只要 0.66 元）
FEE_OPEN = {"cu": 27, "ao": 5.14, "br": 1.63, "sp": 0.95, "rb": 3.06, "hc": 3.22, "ru": 3.03, "ss": 2.02}
FEE_CLOSE = {"cu": 54, "ao": 5.14, "br": 1.63, "sp": 0.02, "rb": 3.06, "hc": 3.22, "ru": 0.03, "ss": 0.02}
COEF_COLS = ["品种", "开盘时间", "场次", "主力", "昨收", "开盘价", "国内跳空%", "外盘涨跌%", "外盘时段", "来源"]


# ---------- 公共的小工具（开盘预估.py 也用） ----------

def bucket(move) -> str:
    if move is None or pd.isna(move):
        return ""
    return next(lab for hi, lab in BINS if abs(move) < hi)


def product_of(symbol: str) -> str:
    """SHFE.ag2612 → ag"""
    return re.match(r"[A-Za-z]+", symbol.split(".")[-1]).group().lower()


def month_no(symbol: str) -> int:
    """SHFE.ag2612 → 2026×12+12，用来算腿1比腿2晚几个月。郑商所三位年月按 202X 年算。"""
    d = re.search(r"(\d{3,4})$", symbol).group(1)
    y, m = (int(d[:2]), int(d[2:])) if len(d) == 4 else (20 + int(d[0]), int(d[1:]))
    return (2000 + y) * 12 + m


def offset_of(s1: str, s2: str) -> int:
    """腿1 比腿2（主力）晚几个月：ag2611 对 ag2612 是 -1。"""
    return month_no(s1) - month_no(s2)


def session_of(open_ts) -> str:
    return "夜盘" if open_ts.hour >= 20 else "日盘"


def num(x):
    """CSV 里的数：空的、写不成数的都当没有。"""
    try:
        v = float(x)
        return None if math.isnan(v) else v
    except (TypeError, ValueError):
        return None


def fmt(x, nd=2, plus=False) -> str:
    v = num(x)
    if v is None:
        return ""
    s = f"{v:+.{nd}f}" if plus else f"{v:.{nd}f}"
    return s.rstrip("0").rstrip(".") if "." in s else s


def direction_of(move) -> str:
    """外盘方向：涨 / 跌；没有数或正好没动是空的。"""
    m = num(move)
    return "" if m is None or abs(m) < 1e-9 else ("涨" if m > 0 else "跌")


def side(change, move) -> str:
    """小草的定义：外盘涨、价差（腿1 − 腿2）比昨收价差往上 = 同向，外盘跌、价差往下也是同向；反过来 = 反向。
    （2026-10-22 起；之前是反过来定的，针_日线基线.csv 的同向%、反向% 两列已对调）
    change = 价差相对昨收价差的变化，move = 外盘涨跌%。"""
    c, m = num(change), num(move)
    if c is None or m is None:
        return ""
    if abs(m) < 1e-9:
        return "外盘没动"
    if abs(c) < 1e-9:
        return "不变"
    return "同向" if (c > 0) == (m > 0) else "反向"


def needle_truth(dev, first, gate) -> str:
    """dev = 针比昨收价差偏几档，first = 首笔真实成交比昨收价差偏几档，gate = 这个品种的门槛（档）。
    真 = 首笔开在针的位置（离针 1 档以内，或比针走得还远）；半真 = 首笔偏出门槛、但比针小；
    假 = 首笔在门槛以内（几乎没动）；反向 = 首笔往针的反方向偏出门槛；没针 = 针本身没偏出门槛，不评。"""
    if gate is None or dev is None or first is None:
        return ""
    d, m = (int(math.copysign(math.floor(abs(x) + 0.5), x)) for x in (dev, first))  # 首笔是半秒成交均价，四舍五入到整档
    if abs(d) < gate:
        return "没针"
    m = m if d > 0 else -m
    if m >= abs(d) - 1:
        return "真"
    if m >= gate:
        return "半真"
    return "反向" if m <= -gate else "假"


def ratio_bin(x) -> str:
    v = num(x)
    return "" if v is None else next(lab for hi, lab in RATIO_BINS if v < hi)


def read_csv(f: Path):
    if not f.exists():
        return None
    return pd.read_csv(f, encoding="utf-8-sig")


def read_ticks(f: Path):
    df = read_csv(f)
    if df is None or not len(df):
        return None
    df["时间"] = pd.to_datetime(df["时间"])
    return df.set_index("时间").sort_index()


# ---------- 系数（开盘预估.py 用） ----------

def load_coef() -> pd.DataFrame:
    df = read_csv(COEF_FILE)
    if df is None:
        return pd.DataFrame(columns=COEF_COLS)
    df["开盘时间"] = pd.to_datetime(df["开盘时间"])
    return df.dropna(subset=["国内跳空%", "外盘涨跌%"]).sort_values("开盘时间")


def fit(coef: pd.DataFrame, product: str, session: str, before, move=None) -> dict | None:
    """国内跳空% ≈ 系数 × 外盘涨跌%（不加截距），用这个品种在 before 之前最近 250 次开盘。
    日盘攒够 20 次之前借夜盘的。返回系数、误差范围（最近 60 次的中间八成）、报警、这一档（同方向、同分档）国内跟了几成。"""
    rows = coef[(coef["品种"] == product) & (coef["开盘时间"] < pd.Timestamp(before))]
    day = rows[rows["场次"] == "日盘"]
    use = day if session == "日盘" and len(day) >= DAY_MIN else rows[rows["场次"] == "夜盘"]
    note = f"日盘只有 {len(day)} 次（攒够 {DAY_MIN} 次才单独算），先借夜盘的系数" if session == "日盘" and use is not day else ""
    use = use.tail(COEF_N)
    if len(use) < 20:
        return None
    x, y = use["外盘涨跌%"].to_numpy(float), use["国内跳空%"].to_numpy(float)
    beta = float(x @ y / (x @ x))
    res = y - beta * x
    band = np.percentile(res[-BAND_N:], [10, 90])
    r20 = res[-ALARM_N:]
    alarms = []
    if np.median(np.abs(r20)) > 2 * np.median(np.abs(res)):
        alarms.append(f"误差变大：最近 {ALARM_N} 次的误差（中位 {np.median(np.abs(r20)):.2f}%）"
                      f"是 {len(res)} 次（{np.median(np.abs(res)):.2f}%）的两倍多，外盘和国内的关系可能变了")
    up, dn = int((r20 > 0).sum()), int((r20 < 0).sum())
    if max(up, dn) >= 15:
        alarms.append(f"偏一边：最近 {ALARM_N} 次有 {max(up, dn)} 次国内开得比预估{'高' if up > dn else '低'}")
    out = {"系数": beta, "次数": len(use), "误差下沿%": float(band[0]), "误差上沿%": float(band[1]),
           "报警": alarms, "说明": note, "起": use["开盘时间"].iloc[0], "止": use["开盘时间"].iloc[-1],
           "来源": "、".join(dict.fromkeys(s.split("（")[0] for s in use["来源"]))}
    b = bucket(move)
    if b and direction_of(move):  # 这一档（外盘同方向、同分档）过去国内跟了几成；外盘涨跌太小时比值没意义，没波动档只看 0.3% 以上的
        all_rows = rows[rows["场次"] == use["场次"].iloc[0]]
        a = all_rows["外盘涨跌%"].abs()
        lo = {"没波动": 0.3, "小": 1, "大": 3, "超大": 6}[b]
        hi = {"没波动": 1, "小": 3, "大": 6, "超大": math.inf}[b]
        s = all_rows[(a >= lo) & (a < hi) & ((all_rows["外盘涨跌%"] > 0) == (move > 0))]
        out["这一档次数"] = len(s)
        out["这一档跟了"] = float(np.median(s["国内跳空%"] / s["外盘涨跌%"])) if len(s) else None
    return out


def needle_base(product: str, offset, dirn: str, bkt: str):
    """两年日线粗估（复盘数据/针_日线基线.csv）：外盘这个方向、这一档，针是同向还是反向、变化多大（针变化占主力昨收的 %，正 = 价差往上）。
    offset 写「全部」= 几个组合合在一起。"""
    base = read_csv(BASE_FILE)
    if base is None or not bkt or not dirn or "外盘方向" not in base:
        return None
    r = base[(base["品种"] == product) & (base["腿1相对主力"].astype(str) == str(offset))
             & (base["外盘方向"] == dirn) & (base["分档"] == bkt)]
    return r.iloc[0].to_dict() if len(r) else None


def tick_history(rec: pd.DataFrame, product: str, offset: int, dirn: str, bkt: str, before) -> dict:
    """开盘复盘记下的 tick（同品种、腿1 相对主力同样几个月、外盘同方向同一档）：针和第一次真实成交各是同向还是反向。"""
    if rec is None or not len(rec) or not bkt or not dirn or "外盘方向" not in rec:
        return {"次数": 0}
    s = rec[(rec["品种"] == product) & (rec["腿1相对主力"] == offset) & (rec["外盘方向"] == dirn) & (rec["分档"] == bkt)
            & (pd.to_datetime(rec["开盘"], format="%Y-%m-%d %H:%M") < pd.Timestamp(before))]
    s = opened(s)
    out = {"次数": len(s)}
    for k in ("针方向", "首笔方向"):
        v = s[k].fillna("")
        out[k] = {x: int((v == x).sum()) for x in ("同向", "反向", "不变")}
    return out


def small(n: int) -> str:
    return "没有样本" if not n else "样本太少，只是线索" if n < 10 else ("样本不多，证据弱" if n < 20 else "")


# ---------- 读一次开盘 ----------

def load_closes(folder: Path, fc) -> dict:
    """{合约: {收盘价, 收盘时间, 到期日}}：收盘价.csv（开盘tick.py 存的），没有就用预估.csv 里记的。"""
    out = {}
    df = read_csv(folder / "收盘价.csv")
    if df is not None:
        for _, r in df.iterrows():
            out[r["合约"]] = {"收盘价": num(r["收盘价"]), "收盘时间": r.get("收盘时间"), "到期日": r.get("到期日"),
                              "最小变动价位": num(r.get("最小变动价位"))}
    if fc is not None:
        for _, r in fc.iterrows():
            for leg in ("腿1", "腿2"):
                out.setdefault(r[leg], {"收盘价": num(r[f"{leg}昨收"]), "收盘时间": r.get("收盘时间"),
                                        "到期日": r.get(f"{leg}到期日")})
    return out


def overseas_saved(folder: Path, product: str, open_ts) -> dict:
    """国内收盘 → 开盘前 7 分钟的外盘涨跌，用开盘文件夹里存的外盘分钟线算；没有分钟线就用外盘涨跌.csv（到开盘前半小时）。"""
    if product not in wp.VARIETIES:
        return {}
    name = f"{wp.VARIETIES[product][0]} {product}"
    tab = read_csv(folder / "外盘涨跌.csv")
    if tab is None:
        return {}
    row = tab[tab["品种"] == name]
    if not len(row) or pd.isna(row.iloc[0].get("从（国内收盘）")):
        return {}
    row = row.iloc[0]
    start = datetime.strptime(f"{open_ts.year}-{row['从（国内收盘）']}", "%Y-%m-%d %H:%M")
    if start > open_ts:  # 跨年：1 月初开盘，收盘在去年 12 月
        start = start.replace(year=start.year - 1)
    end = open_ts - LEAD
    src = row.get("数据源") if isinstance(row.get("数据源"), str) else row.get("东方财富代码")
    raw = read_csv(folder / "外盘分钟线.csv")
    if raw is not None and isinstance(row.get("外盘"), str):
        bars = raw[raw["外盘"] == row["外盘"]].copy()
        bars["时间"] = pd.to_datetime(bars["时间"])
        bars = bars.sort_values("时间")
        a, b = bars[bars["时间"] <= start], bars[bars["时间"] <= end]
        if len(a) and len(b) and a["收"].iloc[-1] and b["收"].iloc[-1]:
            return {"外盘涨跌%": round((b["收"].iloc[-1] / a["收"].iloc[-1] - 1) * 100, 3),
                    "外盘时段": f"{start:%H:%M}→{end:%H:%M}", "外盘来源": f"{row['外盘']}（{src}）"}
    if num(row.get("涨跌幅%")) is not None:
        return {"外盘涨跌%": num(row["涨跌幅%"]), "外盘时段": f"{start:%H:%M}→{end + LEAD - timedelta(minutes=30):%H:%M}",
                "外盘来源": f"{row['外盘']}（{src}；只算到开盘前半小时）"}
    return {}


def analyze(folder: Path, open_ts: datetime, s1: str, s2: str, closes: dict, fc) -> dict:
    p1, p2 = product_of(s1), product_of(s2)
    p = p1 if p1 == p2 else f"{p1}-{p2}"  # 跨品种组合（nr2612-ru2701）单独一类：没有门槛、外盘、系数，不跟同品种的混在一起统计
    rec = {"开盘": f"{open_ts:%Y-%m-%d %H:%M}", "场次": session_of(open_ts), "星期": WEEK[open_ts.weekday()],
           "组合": f"{s1.split('.')[-1]}-{s2.split('.')[-1]}", "品种": p, "腿1": s1.split(".")[-1], "腿2": s2.split(".")[-1],
           "腿1相对主力": offset_of(s1, s2)}
    notes = []
    spread, l1, l2 = (read_ticks(folder / f) for f in (f"价差_{s1}-{s2}.csv", f"{s1}.csv", f"{s2}.csv"))
    c1, c2 = (closes.get(s, {}).get("收盘价") for s in (s1, s2))
    if c1 is not None and c2 is not None:
        rec |= {"腿1昨收": c1, "腿2昨收": c2, "昨收价差": round(c1 - c2, 4)}
    else:
        notes.append("没有昨收价（收盘价.csv、预估.csv 都没有）")
    o = pd.Timestamp(open_ts)

    def auction(leg):
        pre = leg[(leg.index < o) & (leg["成交手数"] > 0)] if leg is not None else []
        return pre.iloc[-1] if len(pre) else None

    a1, a2 = auction(l1), auction(l2)
    if a1 is not None:
        rec |= {"腿1竞价": a1["最新价"], "腿1竞价手数": int(a1["成交手数"])}
        oi, chg = num(a1.get("持仓量")), num(a1.get("持仓变化"))
        if oi is not None and chg is not None:
            rec["腿1竞价持仓变化"] = int(chg)  # 负的 = 竞价里平仓比开仓多
            if oi - chg > 0:
                rec["腿1竞价占持仓‰"] = round(a1["成交手数"] / (oi - chg) * 1000, 2)  # 开盘前的持仓 = 竞价后持仓 − 持仓变化
        if c1 is not None:
            rec["腿1跳空"] = round(a1["最新价"] - c1, 4)
            rec["腿1跳空%"] = round((a1["最新价"] / c1 - 1) * 100, 4)
    if a2 is not None:
        rec |= {"腿2竞价": a2["最新价"], "腿2竞价手数": int(a2["成交手数"])}
        if c2 is not None:
            rec |= {"主力跳空": round(a2["最新价"] - c2, 4), "主力跳空%": round((a2["最新价"] / c2 - 1) * 100, 4)}
    if a1 is not None and a2 is not None:
        rec["针"] = round(a1["最新价"] - a2["最新价"], 4)
        if "昨收价差" in rec:
            rec["针变化"] = round(rec["针"] - rec["昨收价差"], 4)
    else:
        notes.append("、".join(f"{s.split('.')[-1]} 竞价没成交" for s, a in ((s1, a1), (s2, a2)) if a is None))
    if spread is not None:
        after = spread[spread.index >= o]
        first = after[after["均价价差"].notna()]
        if len(first):
            t = first.index[0]
            rec |= {"首笔秒": round((t - o).total_seconds(), 1), "首笔价差": round(first["均价价差"].iloc[0], 4)}
            if "昨收价差" in rec:
                rec["首笔变化"] = round(rec["首笔价差"] - rec["昨收价差"], 4)
        for k in (1, 5, 30, 60):
            w = after[after.index <= o + pd.Timedelta(k, unit="s")]
            if len(w):
                rec[f"价差{k}秒"] = round(w["价差"].iloc[-1], 4)
        m = after[after.index <= o + pd.Timedelta(60, unit="s")]
        if len(m):
            rec |= {"首分钟买价差最低": round(m["买价差"].min(), 4), "首分钟卖价差最高": round(m["卖价差"].max(), 4)}
    tick = closes.get(s1, {}).get("最小变动价位") or TICK.get(p)
    if tick:
        if "针变化" in rec:
            rec["针偏几档"] = round(rec["针变化"] / tick, 1)
        if "首笔变化" in rec:
            rec["首笔偏几档"] = round(rec["首笔变化"] / tick, 1)
    if "针偏几档" in rec and "首笔偏几档" in rec:
        if p in NEEDLE_GATE:
            rec["针真假"] = needle_truth(rec["针偏几档"], rec["首笔偏几档"], NEEDLE_GATE[p])
        else:
            notes.append(f"{p} 还没定针的门槛，不判断真假")
    rec |= overseas_saved(folder, p, open_ts)
    if fc is not None:  # 预估.csv 里记的外盘涨跌就是预估时用的，优先用它
        r = fc[fc["组合"] == rec["组合"]]
        if len(r):
            r = r.iloc[0]
            if num(r.get("外盘涨跌%")) is not None:
                rec |= {"外盘涨跌%": num(r["外盘涨跌%"]), "外盘时段": r.get("外盘时段"), "外盘来源": r.get("外盘来源")}
            rec |= forecast_fields(r)
            if "腿2预估" in rec:
                if "腿2竞价" in rec:
                    rec["预估误差"] = round(rec["腿2竞价"] - rec["腿2预估"], 4)
                    lo, hi = rec.get("预估下沿"), rec.get("预估上沿")
                    if lo is not None and hi is not None:
                        rec["主力在范围内"] = "是" if lo <= rec["腿2竞价"] <= hi else "否"
            lo, hi = num(r.get("针参考下沿")), num(r.get("针参考上沿"))
            if lo is not None and hi is not None and "针" in rec:
                rec["在范围内"] = "是" if lo <= rec["针"] <= hi else "否"
            for k in ("仓储提示", "长期事件"):
                if isinstance(r.get(k), str):
                    rec[k] = r[k]
    move = rec.get("外盘涨跌%")
    rec["分档"], rec["外盘方向"] = bucket(move), direction_of(move)
    rec["针方向"], rec["首笔方向"] = side(rec.get("针变化"), move), side(rec.get("首笔变化"), move)
    if move is None and ("针变化" in rec or "首笔变化" in rec):
        notes.append("没有对应外盘，不判断同向反向" if p not in wp.VARIETIES else "外盘没取到，不判断同向反向")
    # 标签
    rec["周一早盘"] = "是" if rec["场次"] == "日盘" and open_ts.weekday() == 0 else ""
    ct = closes.get(s2, {}).get("收盘时间")
    if isinstance(ct, str) and ct:
        rec["节后"] = "是" if open_ts - datetime.strptime(ct[:16], "%Y-%m-%d %H:%M") > timedelta(days=3.5) else ""
    rec["换月期"] = "是" if 8 <= open_ts.day <= 15 else ""
    ex = closes.get(s1, {}).get("到期日")
    if isinstance(ex, str) and ex:
        rec["腿1离到期天数"] = (datetime.strptime(ex[:10], "%Y-%m-%d").date() - open_ts.date()).days
    rec["备注"] = "；".join(x for x in notes if x)
    return rec


def forecast_fields(r) -> dict:
    """预估.csv 一行里要抄进开盘复盘的：两腿预估开盘价、这次开盘价差的八成范围（预测区间，就是预估页的「针八成落在」）、
    价差往哪走（预测方向）。主力价格的八成范围不写进表格，只给网页用。"""
    out = {}
    if num(r.get("腿2预估")) is not None:
        out |= {"腿1预估": num(r.get("腿1预估")), "腿2预估": num(r["腿2预估"]),
                "预估下沿": num(r.get("腿2下沿")), "预估上沿": num(r.get("腿2上沿"))}
    lo, hi = num(r.get("针参考下沿")), num(r.get("针参考上沿"))
    if lo is not None and hi is not None:
        out["预测区间"] = f"{fmt(lo, 2)} ~ {fmt(hi, 2)}"
    if isinstance(r.get("预测方向"), str) and r["预测方向"]:
        out["预测方向"] = r["预测方向"]
    return out


def forecast_only(open_ts: datetime, fc: pd.DataFrame) -> list:
    """还没开盘（只有 20:53 / 8:53 的预估.csv、还没有 tick）：先把预估记下来，开盘后取完 tick 整行重算。"""
    rows = []
    for _, r in fc.iterrows():
        s1, s2 = str(r["腿1"]), str(r["腿2"])
        rec = {"开盘": f"{open_ts:%Y-%m-%d %H:%M}", "场次": session_of(open_ts), "星期": WEEK[open_ts.weekday()],
               "组合": r["组合"], "品种": r["品种"], "腿1": s1.split(".")[-1], "腿2": s2.split(".")[-1],
               "腿1相对主力": offset_of(s1, s2), "腿1昨收": num(r.get("腿1昨收")), "腿2昨收": num(r.get("腿2昨收")),
               "昨收价差": num(r.get("昨收价差"))}
        if num(r.get("外盘涨跌%")) is not None:
            rec |= {"外盘涨跌%": num(r["外盘涨跌%"]), "外盘时段": r.get("外盘时段"), "外盘来源": r.get("外盘来源")}
        rec |= forecast_fields(r)
        for k in ("仓储提示", "长期事件"):
            if isinstance(r.get(k), str):
                rec[k] = r[k]
        move = rec.get("外盘涨跌%")
        rec["分档"], rec["外盘方向"] = bucket(move), direction_of(move)
        rec["周一早盘"] = "是" if rec["场次"] == "日盘" and open_ts.weekday() == 0 else ""
        rec["换月期"] = "是" if 8 <= open_ts.day <= 15 else ""
        rec["备注"] = "还没开盘：只有开盘前 7 分钟的预估，开盘后取完 tick 整行补全"
        rows.append(rec)
    return rows


def simulate(spread: pd.DataFrame, open_ts, L: float, s: int, tick: float, p: str) -> dict:
    """照标签做一手：挂价 L，s = +1 针往上（挂空价差，卖价差或成交均价差 ≥ L 就成交）、-1 针往下（反过来）。
    成交价按那半秒立刻能成交的价（先看）或成交均价差，比挂价好就按好的算；出场见 SIM_*。"""
    o = pd.Timestamp(open_ts)
    sec = (spread.index - o).total_seconds()
    d = spread[sec > 0]
    t = sec[sec > 0]
    entry, leave = (d["卖价差"], d["买价差"]) if s > 0 else (d["买价差"], d["卖价差"])
    fill = None
    for ti, book, avg in zip(t, entry, d["均价价差"]):
        if ti > SIM_VALID + 1e-9:
            break
        c = [x for x in (num(book), num(avg)) if x is not None and s * (x - L) >= -1e-9]
        if c:
            fill = (ti, c[0])
            break
    out = {"模拟挂价": L}
    if fill is None:
        return out | {"模拟成交": "否"}
    tf, F = fill
    out |= {"模拟成交": "是", "模拟成交秒": round(tf, 1), "模拟成交价差": round(F, 4)}
    side, tgt, last, how = -s, L - s * SIM_TP * tick, None, ""
    for ti, px in zip(t, leave):
        px = num(px)
        if ti <= tf or px is None:
            continue
        last = px
        if (px - tgt) * side >= -1e-9:
            how = f"止盈 {ti:.1f} 秒"
            break
        if ti > tf + SIM_DELAY + 1e-9 and (px - F) * side <= -SIM_SL * tick + 1e-9:
            how = f"止损 {ti:.1f} 秒"
            break
    if last is None:
        return out | {"模拟出场": "成交后没有能平的价"}
    ticks = round((last - F) * side / tick, 1)
    out |= {"模拟出场价差": round(last, 4), "模拟出场": how or f"{t[-1]:.0f} 秒到点（tick 到头）", "模拟档数": ticks}
    if p in MULT:
        gross = round(ticks * tick * MULT[p], 1)
        out["模拟每手毛元"] = gross
        if p in FEE_OPEN:
            fee = round(2 * (FEE_OPEN[p] + FEE_CLOSE[p]), 2)
            out |= {"模拟手续费": fee, "模拟每手净元": round(gross - fee, 1)}
    return out


def sim_cols(folder: Path, open_ts, s1: str, s2: str, rec: dict, tick) -> dict:
    """判 ✅▲ 的照标签模拟一手（模拟成交 是 / 否 / 没有首笔半）；判 ❌ 的也假设照做一次（模拟成交 写「❌假设…」），
    给记分卡数「错过的」和「❌ 省下的」。第 1 步判的 ❌ 没有猜首笔，按混合版用立刻能成交那边补一个。"""
    import 竞价判断 as jd
    k, p = rec.get("59分判断"), rec.get("品种")
    base, needle = num(rec.get("昨收价差")), num(rec.get("针"))
    if k not in ("真", "半真", "假") or p not in NEEDLE_GATE or base is None or needle is None or needle == base or not tick:
        return {}
    s = 1 if needle > base else -1
    half = num(rec.get("首笔半"))
    if half is None and k == "假" and p in jd.MIX:
        e = num(rec.get("立刻能卖" if s > 0 else "立刻能买"))
        if e is not None:
            half = base + round(jd.MIX[p] * round((e - base) / tick, 6), 6) * tick * jd.half_of(p)
    if half is None:
        return {"模拟成交": "没有首笔半"} if k != "假" else {}
    spread = read_ticks(folder / f"价差_{s1}-{s2}.csv")
    if spread is None:
        return {}
    out = simulate(spread, open_ts, round(round(half / tick) * tick, 6), s, tick, p)
    if k == "假":
        out["模拟成交"] = {"是": "❌假设成交", "否": "❌假设没成交"}[out["模拟成交"]]
    return out


def opened(rec: pd.DataFrame) -> pd.DataFrame:
    """真开过盘、有 tick 的行（去掉只有预估的）。"""
    return rec[rec["针"].notna() | rec["首笔价差"].notna()]


def build_records(out: Path = DATA) -> pd.DataFrame:
    import 竞价判断 as jd  # 59 分判断（竞价判断.py 也 import 本文件，放在这里免得互相 import 卡住）
    past = read_csv(REC_FILE)  # 上一次算的结果：回放算首笔区间时用（只看那次以前的开盘）
    rows = []
    for folder in sorted(out.glob("????-??-??_????")):
        pairs = sorted(folder.glob("价差_*.csv"))
        open_ts = datetime.strptime(folder.name, "%Y-%m-%d_%H%M")
        fc = read_csv(folder / "预估.csv")
        if not pairs:
            if fc is not None and len(fc):
                rows += forecast_only(open_ts, fc)
            continue
        closes = load_closes(folder, fc)
        live = jd.read_live(folder)
        for f in pairs:
            s1, s2 = f.stem[len("价差_"):].split("-")
            try:
                rec = analyze(folder, open_ts, s1, s2, closes, fc)
                try:
                    j = live.get(rec["组合"]) or jd.replay_pair(folder, open_ts, s1, s2, closes, past)
                    rec |= jd.review_cols(j, rec.get("首笔价差"), product_of(s1) if product_of(s1) == product_of(s2) else None)
                    if isinstance(j.get("重仓标记"), str) and j["重仓标记"]:
                        rec["重仓标记"] = j["重仓标记"]
                except Exception as e:
                    rec["备注"] = "；".join(x for x in (rec.get("备注"), f"59 分判断算错了：{type(e).__name__} {e}") if x)
                try:
                    rec |= sim_cols(folder, open_ts, s1, s2, rec, closes.get(s1, {}).get("最小变动价位") or TICK.get(rec["品种"]))
                except Exception as e:
                    rec["备注"] = "；".join(x for x in (rec.get("备注"), f"模拟盈亏算错了：{type(e).__name__} {e}") if x)
                rows.append(rec)
            except Exception as e:
                rows.append({"开盘": f"{open_ts:%Y-%m-%d %H:%M}", "组合": f"{s1}-{s2}", "备注": f"算错了：{type(e).__name__} {e}"})
    return pd.DataFrame(rows, columns=REC_COLS + HIDDEN_COLS)


def coef_live(rec: pd.DataFrame) -> pd.DataFrame:
    """每次开盘、每个品种一行：主力（腿2）的跳空 vs 外盘涨跌，给开盘预估算系数用。"""
    rows = []
    ok = rec.dropna(subset=["腿2昨收", "外盘涨跌%"])
    for (o, p), g in ok.groupby(["开盘", "品种"]):
        r = g.iloc[0]
        price = num(r["腿2竞价"])
        if price is None or p not in wp.VARIETIES:
            continue
        rows.append({"品种": p, "开盘时间": o, "场次": r["场次"], "主力": r["腿2"].split(".")[-1], "昨收": r["腿2昨收"],
                     "开盘价": price, "国内跳空%": round((price / r["腿2昨收"] - 1) * 100, 4),
                     "外盘涨跌%": r["外盘涨跌%"], "外盘时段": r["外盘时段"], "来源": f"实盘（{r['外盘来源']}）"})
    return pd.DataFrame(rows, columns=COEF_COLS)


def update_coef(live: pd.DataFrame) -> None:
    """系数历史 = 回测行 + 实盘行；同一次开盘两边都有的，用实盘的。"""
    old = read_csv(COEF_FILE)
    keep = old[old["来源"].astype(str).str.startswith("回测")] if old is not None else pd.DataFrame(columns=COEF_COLS)
    if len(live):
        k = set(zip(live["品种"], live["开盘时间"]))
        keep = keep[[(p, t) not in k for p, t in zip(keep["品种"], keep["开盘时间"])]]
    df = pd.concat([keep, live], ignore_index=True) if len(live) else keep
    df = df.sort_values(["品种", "开盘时间"])
    REVIEW.mkdir(exist_ok=True)
    df.to_csv(COEF_FILE, index=False, encoding="utf-8-sig")


# ---------- 记分卡 ----------

def count_sides(s: pd.Series) -> str:
    v = s.fillna("")
    same, opp = int((v == "同向").sum()), int((v == "反向").sum())
    n = same + opp + int((v == "不变").sum())
    if not n:
        return "没有样本"
    return f"同向 {same}、反向 {opp}" + (f"、不变 {n - same - opp}" if n - same - opp else "") + f"（{n} 次）"


def daily_sides(base, dirn: str, bkt: str) -> str:
    """两年日线：白银、原油、燃油几个组合合在一起，针同向、反向各占多少（按组合数加权）。"""
    if base is None or "外盘方向" not in base:
        return ""
    d = base[(base["外盘方向"] == dirn) & (base["分档"] == bkt) & (base["品种"].isin(["ag", "sc", "fu"]))
             & (base["腿1相对主力"].astype(str) == "全部")]
    if not len(d):
        return "没有样本"
    w = d["次数"]
    return (f"同向 {np.average(d['同向%'], weights=w):.0f}%、反向 {np.average(d['反向%'], weights=w):.0f}%"
            f"（{int(d['开盘次数'].sum())} 次开盘）")


def per_opening(rec: pd.DataFrame) -> pd.DataFrame:
    """每次开盘、每个品种一行（涨跌幅统计用，几个组合不重复算）。"""
    return rec.dropna(subset=["外盘涨跌%"]).drop_duplicates(["开盘", "品种"])


def rule_team(rec: pd.DataFrame) -> str:
    """规则一（队里）：大涨大跌时，第一次真实成交同向。外盘涨、跌分开，每档列出外盘和主力的涨跌幅。"""
    rows = []
    base = read_csv(BASE_FILE)
    one = per_opening(rec)
    for dirn in DIRS:
        for b in BUCKETS:
            s = rec[(rec["外盘方向"] == dirn) & (rec["分档"] == b)]
            o = one[(one["外盘方向"] == dirn) & (one["分档"] == b)]
            n_open = s["开盘"].nunique()
            rows.append([f"外盘{dirn}", b, f"{n_open} 次开盘", fmt(o["外盘涨跌%"].median(), 2, plus=True),
                         fmt(o["主力跳空%"].median(), 2, plus=True), count_sides(s["首笔方向"]), count_sides(s["针方向"]),
                         small(n_open), daily_sides(base, dirn, b)])
    return table(["外盘", "分档", "涉及", "外盘涨跌中位%", "主力跳空中位%", "第一次真实成交", "针（竞价）", "样本",
                  "两年日线：针（白银、原油、燃油合计，粗估）"], rows)


def follow(g: pd.DataFrame) -> str:
    """国内跟了外盘几成：主力跳空% ÷ 外盘涨跌% 的中位数（外盘不到 0.3% 的不算，比值没意义）。"""
    g = g[g["外盘涨跌%"].abs() >= 0.3]
    return fmt((g["国内跳空%"] / g["外盘涨跌%"]).median(), 2) if len(g) else ""


def move_stats(rec: pd.DataFrame) -> str:
    """涨跌幅统计：每个品种，外盘涨、跌分开，每档外盘和国内主力各涨跌多少、国内跟了几成；旁边是两年回测（系数历史.csv 的回测行）。"""
    one = per_opening(rec).copy()
    one["国内跳空%"] = one["主力跳空%"]
    coef = load_coef()
    back = coef[coef["来源"].astype(str).str.startswith("回测")].copy()
    back["外盘方向"] = back["外盘涨跌%"].map(direction_of)
    back["分档"] = back["外盘涨跌%"].map(bucket)
    rows = []
    for p in sorted(one["品种"].unique()):
        for dirn in DIRS:
            for b in BUCKETS:
                g = one[(one["品种"] == p) & (one["外盘方向"] == dirn) & (one["分档"] == b)]
                if not len(g):
                    continue
                h = back[(back["品种"] == p) & (back["外盘方向"] == dirn) & (back["分档"] == b)]
                old = (f"{len(h)} 次，外盘中位 {fmt(h['外盘涨跌%'].median(), 2, plus=True)}%、"
                       f"国内中位 {fmt(h['国内跳空%'].median(), 2, plus=True)}%、跟了 {follow(h) or '—'}") if len(h) else "没有"
                rows.append([p, f"外盘{dirn}", b, f"{len(g)} 次", fmt(g["外盘涨跌%"].median(), 2, plus=True),
                             fmt(g["主力跳空%"].median(), 2, plus=True), follow(g) or "—", small(len(g)), old])
    if not rows:
        return "<p>还没有带外盘涨跌的开盘。</p>"
    return table(["品种", "外盘", "分档", "实盘次数", "外盘涨跌中位%", "主力跳空中位%", "国内跟了几成", "样本",
                  "两年回测（夜盘，外盘到 21:00）"], rows)


def rule_fu(rec: pd.DataFrame) -> str:
    """规则二：燃油 23 点收盘后原油大跌，第二天燃油开盘有大动静。看日盘燃油：外盘（23:00→8:53）涨跌多少，跌和涨分开。"""
    s = rec[(rec["品种"] == "fu") & (rec["场次"] == "日盘")].dropna(subset=["外盘涨跌%"])
    rows = []
    x = s["外盘涨跌%"]
    for lab, m in (("布伦特跌 3% 以上", x <= -3), ("跌 1%–3%", (x > -3) & (x <= -1)), ("1% 以内", (x > -1) & (x < 1)),
                   ("涨 1%–3%", (x >= 1) & (x < 3)), ("涨 3% 以上", x >= 3)):
        g = s[m]
        n = g["开盘"].nunique()
        rows.append([lab, f"{n} 次开盘", fmt(g.groupby("开盘")["主力跳空%"].first().abs().median(), 2),
                     fmt(g["针变化"].abs().median(), 1), fmt(g["首笔变化"].abs().median(), 1), small(n) if n else "没有样本"])
    return table(["日盘燃油", "次数", "主力跳空中位（%，绝对值）", "针变化中位（元，绝对值）", "第一次成交变化中位（元）", "样本"], rows)


def rule_monday(rec: pd.DataFrame) -> str:
    """规则三：周一早上开盘大概率有东西。周一 9:00 对比其他日盘开盘。"""
    s = rec[rec["场次"] == "日盘"]
    rows = []
    for lab, m in (("周一早盘", s["周一早盘"] == "是"), ("其他日盘", s["周一早盘"] != "是")):
        g = s[m]
        n = g["开盘"].nunique()
        jump = g.groupby(["开盘", "品种"])["主力跳空%"].first().abs()
        rows.append([lab, f"{n} 次开盘", fmt(jump.median(), 2), fmt(g["针变化"].abs().median(), 1),
                     fmt(g["首笔变化"].abs().median(), 1), small(n) if n else "没有样本"])
    return table(["", "次数", "主力跳空中位（%，绝对值）", "针变化中位（元，绝对值）", "第一次成交变化中位（元）", "样本"], rows)


def rule_ratio(rec: pd.DataFrame) -> str:
    """规则四：腿1竞价占持仓越高，针越可能真或半真。每个品种、每档数真 / 半真 / 假 / 反向（没针的另列，不算比例）。"""
    s = rec[rec["针真假"].isin(TRUTHS)].copy()
    if not len(s):
        return "<p>还没有判断过真假的开盘。</p>"
    s["占持仓档"] = s["腿1竞价占持仓‰"].map(ratio_bin)
    rows = []
    for p in ["全部"] + sorted(s["品种"].unique()):
        g = s if p == "全部" else s[s["品种"] == p]
        for _, lab in RATIO_BINS:
            h = g[g["占持仓档"] == lab]
            if not len(h):
                continue
            c = {k: int((h["针真假"] == k).sum()) for k in TRUTHS}
            judged = c["真"] + c["半真"] + c["假"] + c["反向"]
            good = f"{(c['真'] + c['半真']) / judged * 100:.0f}%" if judged else ""
            dec = int((h[h["针真假"].isin(["假", "反向"])]["腿1竞价持仓变化"] < 0).sum())
            rows.append([p, lab, c["真"], c["半真"], c["假"], c["反向"], good, dec, c["没针"],
                         small(h["开盘"].nunique()) if judged else ""])
    return table(["品种", "腿1竞价占持仓", "真", "半真", "假", "反向", "真+半真占比", "假/反向里竞价后持仓减少", "没针（不评）",
                  "样本"], rows)


def rule_auction(rec: pd.DataFrame) -> str:
    """规则五：59 分判断准不准。判成 ✅ ▲ ❌ 的，实际开成 真 / 半真 / 假 / 反向 各几根；首笔落没落在预计区间里。实盘、回放分开。"""
    import 竞价判断 as jd
    s = rec[rec["59分判断"].isin(["真", "半真", "假"])]
    rows = []
    for src in ("实盘", "回放"):
        for k in ("真", "半真", "假"):
            g = s[(s["判断来源"] == src) & (s["59分判断"] == k)]
            if not len(g):
                continue
            c = {x: int((g["针真假"] == x).sum()) for x in ("真", "半真", "假", "反向")}
            n = sum(c.values())
            right = c[k] + (c["反向"] if k == "假" else 0)  # ❌ 说的是「开不到」，实际反向也算对
            fill = c["真"] + c["半真"] if k != "假" else c["假"] + c["反向"]  # 能不能开到说对了没有
            box = g["首笔在区间里"].isin(["是", "否"])
            rows.append([src, f"{jd.SYMBOL[k]} {k}", n, c["真"], c["半真"], c["假"], c["反向"],
                         f"{right / n * 100:.0f}%" if n else "", f"{fill / n * 100:.0f}%" if n else "",
                         f"{int((g['首笔在区间里'] == '是').sum())} / {int(box.sum())}", small(g["开盘"].nunique())])
    if not rows:
        return "<p>还没有判成 ✅ ▲ ❌ 的针。</p>"
    box = rec["首笔在区间里"].isin(["是", "否"])
    allin = (f"<p class='note'>首笔区间（所有有区间的组合，含没针的）：{int((rec['首笔在区间里'] == '是').sum())} / {int(box.sum())} "
             f"落在里面。</p>") if box.any() else ""
    return table(["来源", "判成", "根数", "实际真", "半真", "假", "反向", "说对了", "能不能开到说对了", "首笔在区间里", "样本"],
                 rows) + allin + versus_v1(s, jd.MIX)


def versus_v1(s: pd.DataFrame, mixed) -> str:
    """换成混合版的品种（sc）：同一批针，现在的判断和 v1 各说对几根。"""
    def score(g, col):
        k, act = g[col], g["针真假"]
        right = (k == act) | ((k == "假") & (act == "反向"))  # ❌ 说的是「开不到」，实际反向也算对
        fill = k.isin(["真", "半真"]) == act.isin(["真", "半真"])
        return int(right.sum()), int(fill.sum())

    rows = []
    for p in mixed:
        g = s[(s["品种"] == p) & s["v1判断"].isin(["真", "半真", "假"]) & s["针真假"].isin(["真", "半真", "假", "反向"])]
        for src in ("实盘", "回放"):
            h = g[g["判断来源"] == src]
            if len(h):
                (a, b), (c, d) = score(h, "59分判断"), score(h, "v1判断")
                n = len(h)
                rows.append([p, src, n, f"{a} / {n}", f"{b} / {n}", f"{c} / {n}", f"{d} / {n}",
                             int((h["59分判断"] != h["v1判断"]).sum()), small(h["开盘"].nunique())])
    if not rows:
        return ""
    return ("<h4>换成混合版的品种：同一批针 v1 会怎样</h4>"
            "<p class='note'>sc 2026-10-24 起第 2 步换成混合版（两年 L2 回放：153 根针三档说对 59% → 71%，开不开得到 66% → 76%；10-24 晚放宽 ❌ 后是 66%、73%，换来真针判成 ❌ 7 → 3 根）。"
            "10-24 以前的实盘是 v1 贴的，两列一样。</p>"
            + table(["品种", "来源", "根数", "现在说对了", "能不能开到说对了", "v1 说对了", "v1 能不能开到说对了", "判得不一样", "样本"],
                    rows))


def sim_stats(g: pd.DataFrame) -> dict:
    """一组模拟（模拟成交 = 是 / ❌假设成交 的算成交）：次数、胜率、每次平均档数、毛元、扣费后元。"""
    f = g[g["模拟成交"].isin(["是", "❌假设成交"])]
    x = pd.to_numeric(f["模拟档数"], errors="coerce").dropna()
    gross = pd.to_numeric(f["模拟每手毛元"], errors="coerce").dropna()
    net = pd.to_numeric(f["模拟每手净元"], errors="coerce").dropna()
    return {"成交": len(x), "赢": int((x > 0).sum()), "胜率": f"{(x > 0).mean() * 100:.0f}%" if len(x) else "",
            "平均档": fmt(x.mean(), 1, plus=True) if len(x) else "", "合计档": fmt(x.sum(), 0, plus=True) if len(x) else "",
            "平均毛元": fmt(gross.mean(), 0, plus=True) if len(gross) else "",
            "平均净元": f"{fmt(net.mean(), 0, plus=True)}（{len(net)} 次）" if len(net) else "—"}


def rule_money(rec: pd.DataFrame) -> str:
    """规则六：照浮层标签做赚不赚钱。判 ✅▲ 的按品种、符号、重 / 不重、实盘 / 回放累计；判 ❌ 但实际开到的另列「错过的」。"""
    s = rec[rec["59分判断"].isin(["真", "半真"])].copy()
    if not len(s):
        return "<p>还没有判成 ✅ ▲ 的针。</p>"
    s["符号"] = s["59分判断"].map({"真": "✅", "半真": "▲"})
    s["重否"] = np.where(s["重仓标记"].astype(str).str.startswith("重"), "重", "不重")
    rows = []
    dims = [("全部", None), ("来源", "判断来源"), ("符号", "符号"), ("重 / 不重", "重否"), ("品种", "品种")]
    for dim, col in dims:
        for lab, g in ([("全部", s)] if col is None else sorted(s.groupby(col), key=lambda kv: str(kv[0]))):
            st = sim_stats(g)
            rows.append([dim, lab, len(g), st["成交"], int((g["模拟成交"] == "否").sum()), int((g["模拟成交"] == "没有首笔半").sum()),
                         st["赢"], st["胜率"], colored(st["平均档"], 1) if st["平均档"] else "", st["合计档"], st["平均毛元"],
                         st["平均净元"], g["开盘"].nunique(), small(g["开盘"].nunique())])
    out = table(["看", "分组", "判 ✅▲", "成交", "没成交", "没有首笔半", "赢", "胜率", "每次平均档", "合计档", "每次平均毛元",
                 "扣费后每次平均元", "涉及开盘", "样本"], rows)
    # 判 ❌ 的：实际开到了（真 / 半真）= 错过的；实际假 / 反向 = 判对了，假设照做会怎样
    x = rec[rec["59分判断"] == "假"]
    miss = x[x["针真假"].isin(["真", "半真"])]
    right = x[x["针真假"].isin(["假", "反向"])]
    summ = []
    for lab, g in (("❌ 但实际真 / 半真（错过的）", miss), ("❌ 实际假 / 反向（判对了）", right)):
        st = sim_stats(g)
        summ.append([lab, len(g), st["成交"], st["赢"], colored(st["平均档"], 1) if st["平均档"] else "", st["合计档"],
                     st["平均毛元"], g["开盘"].nunique(), small(g["开盘"].nunique()) if len(g) else "没有样本"])
    out += ("<h4>判 ❌ 的：假设照样在首笔半挂一手</h4>"
            + table(["", "根数", "假设成交", "赢", "每次平均档", "合计档", "每次平均毛元", "涉及开盘", "样本"], summ))
    if len(miss):
        out += table(["开盘", "组合", "来源", "实际", "针偏几档", "首笔偏几档", "假设挂价", "假设成交价差", "出场", "档数", "毛元"],
                     [[r["开盘"], r["组合"], r["判断来源"], r["针真假"], fmt(r["针偏几档"], 1), fmt(r["首笔偏几档"], 1),
                       fmt(r["模拟挂价"], 4), fmt(r["模拟成交价差"], 4),
                       r["模拟出场"] if isinstance(r["模拟出场"], str) else ("没成交" if r["模拟成交"] == "❌假设没成交" else ""),
                       colored(r["模拟档数"], 1), fmt(r["模拟每手毛元"], 0, plus=True)]
                      for _, r in miss.sort_values("开盘", ascending=False).iterrows()])
    last = s[s["开盘"] == rec["开盘"].max()]
    if len(last):
        out += (f"<h4>最近一次开盘（{esc(rec['开盘'].max())}）判 ✅▲ 的</h4>"
                + table(["组合", "判", "重仓标记", "挂价（首笔半）", "成交", "成交秒", "成交价差", "出场", "出场价差", "档数", "毛元", "扣费后元"],
                        [[r["组合"], r["符号"], r["重仓标记"] if isinstance(r["重仓标记"], str) else "", fmt(r["模拟挂价"], 4),
                          r["模拟成交"], fmt(r["模拟成交秒"], 1), fmt(r["模拟成交价差"], 4),
                          r["模拟出场"] if isinstance(r["模拟出场"], str) else "", fmt(r["模拟出场价差"], 4),
                          colored(r["模拟档数"], 1), fmt(r["模拟每手毛元"], 0, plus=True), fmt(r["模拟每手净元"], 0, plus=True)]
                         for _, r in last.iterrows()]))
    return out


def accuracy(rec: pd.DataFrame) -> str:
    s = rec.dropna(subset=["预估误差"])
    rows = []
    for (p, sess), g in s.groupby(["品种", "场次"]):
        g = g.groupby("开盘").first()
        err, gap = g["预估误差"].abs(), g["主力跳空"].abs()
        inside = (g["主力在范围内"] == "是").mean() * 100
        rows.append([p, sess, f"{len(g)} 次", fmt(err.mean(), 1), fmt(gap.mean(), 1), f"{inside:.0f}%", small(len(g))])
    if not rows:
        return "<p>还没有预估过的开盘（开盘预估.py 每天 20:53、8:53 自动跑，下一次开盘后这里就有了）。</p>"
    return table(["品种", "场次", "次数", "预估平均差（元）", "不看外盘平均差（元）", "落在八成范围里", "样本"], rows)


# ---------- 网页 ----------

CSS = """
:root{--bg:#f7f6f2;--fg:#1d2433;--muted:#667085;--line:#e3e0d8;--card:#fff;--accent:#1f4e79;--warn:#b54708;--warnbg:#fff4e5;--up:#c0392b;--down:#1e7d4f}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){--bg:#14171c;--fg:#e6e8eb;--muted:#9aa3ae;--line:#2c313a;--card:#1b1f26;--accent:#7fb3e6;--warn:#f5a623;--warnbg:#2b2214;--up:#ff6b5e;--down:#4cc38a}}
:root[data-theme="dark"]{--bg:#14171c;--fg:#e6e8eb;--muted:#9aa3ae;--line:#2c313a;--card:#1b1f26;--accent:#7fb3e6;--warn:#f5a623;--warnbg:#2b2214;--up:#ff6b5e;--down:#4cc38a}
body{margin:0;background:var(--bg);color:var(--fg);font:15px/1.65 "Segoe UI","Microsoft YaHei",sans-serif}
main{max-width:1180px;margin:0 auto;padding:20px 16px 60px}
h1{font-size:22px;margin:0 0 4px}h2{font-size:18px;margin:30px 0 8px;padding-left:8px;border-left:4px solid var(--accent)}
h3{font-size:16px;margin:4px 0 6px}h4{font-size:14px;margin:12px 0 2px;color:var(--muted)}
ul.list{margin:2px 0 6px;padding-left:20px;font-size:14px}ul.list li{margin:2px 0}
.meta,.note{color:var(--muted);font-size:13px}.note{margin:4px 0}
.card{background:var(--card);border:1px solid var(--line);border-radius:8px;padding:12px 14px;margin:12px 0}
.warn{background:var(--warnbg);border-left:4px solid var(--warn);padding:6px 12px;border-radius:4px;margin:8px 0}
.wrap{overflow-x:auto}table{border-collapse:collapse;width:100%;font-size:14px;margin:6px 0}
th,td{border-bottom:1px solid var(--line);padding:5px 8px;text-align:right;white-space:nowrap}
th{color:var(--muted);font-weight:600}td:first-child,th:first-child{text-align:left}
.big{font-size:19px;font-weight:700}.up{color:var(--up)}.down{color:var(--down)}
a{color:var(--accent)}
"""


def esc(x) -> str:
    return html.escape("" if x is None or (isinstance(x, float) and math.isnan(x)) else str(x))


def table(head: list, rows: list) -> str:
    h = "".join(f"<th>{esc(x)}</th>" for x in head)
    b = "".join("<tr>" + "".join(f"<td>{c if isinstance(c, Raw) else esc(c)}</td>" for c in r) + "</tr>" for r in rows)
    return f'<div class="wrap"><table><thead><tr>{h}</tr></thead><tbody>{b}</tbody></table></div>'


class Raw(str):
    """已经是 HTML 的格子，不再转义。"""


def page(title: str, body: str) -> str:
    return (f'<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">'
            f'<meta name="viewport" content="width=device-width,initial-scale=1"><title>{esc(title)}</title>'
            f"<style>{CSS}</style></head><body><main>{body}</main></body></html>")


def colored(x, nd=1) -> Raw:
    v = num(x)
    if v is None:
        return Raw("")
    return Raw(f'<span class="{"up" if v > 0 else "down" if v < 0 else ""}">{fmt(v, nd, plus=True)}</span>')


def move_cell(r) -> str:
    """外盘涨跌 + 方向和分档，如「-1.57% 跌·小」。"""
    if num(r["外盘涨跌%"]) is None:
        return "没有外盘"
    return f"{fmt(r['外盘涨跌%'], 2, plus=True)}% {r['外盘方向']}·{r['分档']}"


def forecast_link(opening: str) -> str:
    """这次开盘的预估网页（预估/2026-10-22_2100.html），没有就空着。"""
    f = FORECAST_DIR / f"{opening.replace(' ', '_').replace(':', '')}.html"
    return f"../{FORECAST_DIR.name}/{f.name}" if f.exists() else ""


def tags(r) -> str:
    """明细里的标签：周一早盘、节后、换月期、腿1 快到期，加上预估时记下的仓储提示（几条【高】【中】）和长期事件。"""
    out = [k for k in ("周一早盘", "节后", "换月期") if r[k] == "是"]
    if num(r["腿1离到期天数"]) is not None and r["腿1离到期天数"] <= 10:
        out.append(f"腿1离到期{int(r['腿1离到期天数'])}天")
    w = r["仓储提示"] if isinstance(r["仓储提示"], str) else ""
    lv = "".join(f"【{k}】{w.count(f'【{k}】')} 条" for k in ("高", "中") if f"【{k}】" in w)
    if lv:
        out.append("仓储" + lv)
    if isinstance(r["长期事件"], str) and r["长期事件"]:
        out.append(r["长期事件"])
    return "、".join(out)


def latest_table(g: pd.DataFrame) -> str:
    rows = []
    for _, r in g.iterrows():
        rows.append([r["组合"], move_cell(r), fmt(r["昨收价差"], 2), fmt(r["针"], 2), colored(r["针变化"], 2), r["针方向"],
                     fmt(r["首笔价差"], 2), colored(r["首笔变化"], 2), r["首笔方向"], fmt(r["首笔秒"], 1),
                     fmt(r["价差5秒"], 2), fmt(r["价差60秒"], 2), fmt(r["首分钟买价差最低"], 2), fmt(r["首分钟卖价差最高"], 2),
                     colored(r["腿1跳空"], 2), colored(r["主力跳空"], 2), r["针真假"], judged(r),
                     r["预计首笔区间"] if isinstance(r["预计首笔区间"], str) else "", r["首笔在区间里"],
                     fmt(r["腿1竞价占持仓‰"], 2), r["备注"]])
    return table(["组合", "外盘", "昨收价差", "针", "针变化", "针方向", "第一次真实成交", "变化", "方向", "开盘后几秒",
                  "5 秒", "60 秒", "第一分钟买价差最低", "卖价差最高", "腿1跳空", "主力跳空", "针真假", "59 分判断",
                  "预计首笔区间", "在区间里", "腿1竞价占持仓‰", "备注"], rows)


def judged(r) -> str:
    """59 分判断，带符号：「✅ 真」；回放补算的后面标「（回放）」。"""
    k = r["59分判断"] if isinstance(r["59分判断"], str) else ""
    if not k:
        return ""
    sym = {"真": "✅", "半真": "▲", "假": "❌", "没针": "○", "不判断": "？"}.get(k, "")
    return f"{sym} {k}" + ("（回放）" if r["判断来源"] == "回放" else "")


def latest_page() -> Path | None:
    """最新的复盘网页（开盘预估页链过来用）。"""
    pages = sorted(PAGE_DIR.glob("????-??-??_????.html"))
    return pages[-1] if pages else None


def write_page(rec: pd.DataFrame) -> Path | None:
    rec = opened(rec)
    if not len(rec):
        return None  # 网页只看开过盘的；还没开盘的预估在 预估/ 的网页里
    pages = sorted(FORECAST_DIR.glob("????-??-??_????.html"))
    link = f"<a href='../{FORECAST_DIR.name}/{pages[-1].name}'>最新的开盘预估</a>（全部在 {FORECAST_DIR.name}/ 文件夹）" if pages else "还没有开盘预估"
    body = [f"<h1>开盘复盘</h1><p class='meta'>更新于 {datetime.now():%Y-%m-%d %H:%M}，"
            f"共 {rec['开盘'].nunique()} 次开盘、{len(rec)} 行（数据/ 里有 tick 的开盘）。"
            f"{link} · 明细在 复盘数据/开盘复盘.csv</p>"
            "<p class='note'>同向 / 反向（小草定的）：外盘涨、价差（腿1 − 腿2，次主力 − 主力）比昨收价差往上 = 同向，"
            "外盘跌、价差往下也是同向；反过来是反向。针和第一次真实成交各判断一次，外盘涨、跌分开统计。</p>"]
    if len(rec):
        last = rec["开盘"].max()
        g = rec[rec["开盘"] == last]
        ov = g.dropna(subset=["外盘涨跌%"]).drop_duplicates("品种")
        ovs = "；".join(f"{r['品种']} 外盘 {move_cell(r)}（{r['外盘时段']}），主力跳空 {fmt(r['主力跳空%'], 2, plus=True)}%"
                       for _, r in ov.iterrows())
        fl = forecast_link(last)
        more = f" · <a href='{fl}'>这次的预估</a>" if fl else ""
        body.append(f"<h2>最近一次开盘：{last}（{g['场次'].iloc[0]}，星期{g['星期'].iloc[0]}）</h2>"
                    f"<p class='note'>{esc(ovs or '这次没有外盘数据')}{more}</p>"
                    "<p class='note'>针 = 竞价那张快照两腿的价差；变化都是相对昨收价差。"
                    "第一次真实成交 = 开盘后两腿第一次在同一个半秒都有成交，那半秒的均价价差。"
                    "针真假（小草定的）：真 = 首笔开在针的位置；半真 = 首笔比昨收偏出门槛、但比针小；假 = 首笔几乎没动；"
                    "没针 = 针本身没偏出门槛。</p>" + latest_table(g))
        fc = g.dropna(subset=["腿2预估"]).drop_duplicates("品种")
        if len(fc):
            out = ("<p class='note'>超出八成范围的，先看这次的预估页里当晚的新闻、仓储提示有没有线索。</p>"
                   if (fc["主力在范围内"] == "否").any() else "")
            body.append(f"<h3>预估对答案</h3>{out}" + table(
                ["品种", "主力", "外盘", "预估", "八成范围", "实际竞价", "差", "在范围里"],
                [[r["品种"], r["腿2"].split(".")[-1], move_cell(r), fmt(r["腿2预估"], 2),
                  f"{fmt(r['预估下沿'], 2)} ~ {fmt(r['预估上沿'], 2)}", fmt(r["腿2竞价"], 2), colored(r["预估误差"], 2),
                  r["主力在范围内"]] for _, r in fc.iterrows()]))
    body.append("<h2>记分卡</h2><p class='note'>同一次开盘的几个组合各算一次，所以「同向 / 反向」的次数比开盘次数多；看「涉及几次开盘」。"
                "不到 10 次开盘只是线索，不到 20 次证据弱。</p>")
    body.append("<div class='card'><h3>规则一（队里）：大涨大跌时，第一次真实成交同向</h3>"
                "<p class='note'>小草：同向基本都是超大波动，平常反向多一点。分档看外盘（国内收盘 → 开盘前 7 分钟），涨、跌分开。"
                "两年日线那一列用日线开盘价当竞价价（不一定是），外盘算到 21:00，只当参考；第一次真实成交日线看不到。</p>"
                + rule_team(rec) + "</div>")
    body.append("<div class='card'><h3>规则二：燃油晚上 11 点收盘后原油大跌，第二天燃油开盘有大动静</h3>"
                "<p class='note'>看日盘燃油开盘：布伦特从 23:00 到 8:53 涨跌多少，对比燃油主力跳空和价差变化；涨的也列出来对照。</p>"
                + rule_fu(rec) + "</div>")
    body.append("<div class='card'><h3>规则三：周一大概率开盘会有东西</h3>"
                "<p class='note'>周一早上 9:00 开盘（周末之后）对比其他日盘开盘。</p>" + rule_monday(rec) + "</div>")
    gates = "、".join(f"{k} {v}" for k, v in NEEDLE_GATE.items())
    body.append("<div class='card'><h3>规则四：腿1竞价成交占持仓越多，针越可能真</h3>"
                "<p class='note'>看开盘第一笔真实成交（小草定的）：真 = 开在针的位置（离针 1 档以内或更远）；"
                "半真 = 比昨收价差偏出门槛、但比针小（照样能接到）；假 = 门槛以内、几乎没动；反向 = 往针的反方向偏出门槛；"
                f"没针 = 针本身没偏出门槛，不评。门槛（档）：{gates}。"
                "腿1竞价占持仓 = 腿1竞价成交手数 ÷ 开盘前腿1持仓（千分之几），不同品种一手大小不同，所以比比例不比手数；"
                "分档的线（0.7‰、1.5‰）是 10-22 从 3 次开盘看出来的临时线。"
                "竞价后持仓减少 = 竞价里平仓比开仓多，10-22 夜盘 fu2702 竞价 3054 手却是假针，持仓少了 2994 手。</p>"
                + rule_ratio(rec) + "</div>")
    body.append("<div class='card'><h3>规则五：59 分判断准不准（✅ ▲ ❌）</h3>"
                "<p class='note'>竞价判断.py 在 20:59 / 8:59 竞价一出结果就判断，贴在博易云的格子上：✅ 大概率整根开到针附近（真）、"
                "▲ 大概率开到一部分（半真）、❌ 大概率假。这里对答案（针真假，小草定的）。说对了 = ✅ 实际真、▲ 实际半真、❌ 实际假或反向；"
                "能不能开到说对了 = ✅ ▲ 实际真或半真、❌ 实际假或反向。实盘 = 当时真贴出来的；回放 = 那次没跑竞价判断、用存下来的 tick 补算的"
                "（v1 是 10-23 看着 10-21 到 10-22 的针定的，这几次的回放是样本内，会偏乐观）。"
                "sc 10-24 起换成混合版、首笔区间换成最近 150 次开盘的八成范围（见竞价判断的网页），回放按现在的规则重算。</p>"
                + rule_auction(rec) + "</div>")
    fees = "、".join(f"{p} 一组 {fmt(2 * (FEE_OPEN[p] + FEE_CLOSE[p]))} 元（开 {FEE_OPEN[p]}、平今 {FEE_CLOSE[p]} 元/手/腿）" for p in FEE_OPEN)
    body.append("<div class='card'><h3>规则六：照浮层标签做，赚不赚钱（模拟，每手）</h3>"
                "<p class='note'>判得准不等于赚钱：真针判成 ❌ 就少赚，▲ 里的假针多半挂不进、不亏。这里按钱记分，以后改参数看这张表的钱有没有变多。"
                f"做法（照抄两年回放）：判 ✅▲ 的在首笔半挂一手，开盘后 {SIM_VALID} 秒内那半秒的成交均价差或立刻能成交的价比挂价好就成交，"
                "成交价按那个更好的价算（不按挂价）；"
                f"止盈 = 挂价往有利方向 {SIM_TP} 档，止损 = 比成交价亏 {SIM_SL} 档，成交后 {SIM_DELAY} 秒内不止损，都按当时立刻能平的价平；"
                "tick 只到开盘后 60 秒，到时按最后能平的价算（两年回放最多拿 300 秒）。全市场的量都算成自己的，是上限；腿2 当成和腿1 同时成交。"
                "「档」= 最小变动价位，毛元 = 档 × 最小变动价位 × 合约乘数（不同品种一手大小差很多，元数平均主要看 sc）。"
                f"手续费只有 {fees or '（还没填）'}，是他的成交记录算的；别的品种还没数，只算毛的（「扣费后」写 —）。"
                "首笔半按现在的系数算（ag 10-31 以前浮层上是挂一半，现在是三分之一）；10-23 实盘还是 v1，没有首笔半，不算。"
                "实盘才十来次开盘，结论看「样本」那一列。</p>"
                + rule_money(rec) + "</div>")
    body.append("<h2>涨跌幅统计</h2><p class='note'>每次开盘、每个品种算一次。外盘涨跌 = 国内收盘 → 开盘前 7 分钟；"
                "主力跳空 = 主力（腿2）竞价价对昨收；国内跟了几成 = 主力跳空 ÷ 外盘涨跌的中位数（外盘不到 0.3% 的不算）。"
                "右边的两年回测是 2024-05 起的夜盘开盘（日线开盘价，外盘算到 21:00），只当对照。</p>" + move_stats(rec))
    body.append("<h2>预估准不准</h2><p class='note'>预估 = 昨收 ×（1 + 系数 × 外盘涨跌），对的是主力（腿2）的竞价价。"
                "「不看外盘」= 直接拿昨收当开盘价的差，预估要比它小才有用。</p>" + accuracy(rec))
    if len(rec):
        recent = rec[rec["开盘"].isin(sorted(rec["开盘"].unique())[-20:])].sort_values(["开盘", "组合"], ascending=[False, True])
        body.append("<h2>最近 20 次开盘明细</h2>" + table(
            ["开盘", "场次", "组合", "外盘", "昨收价差", "针变化", "针方向", "第一次成交变化", "方向", "腿1跳空", "主力跳空",
             "针真假", "59 分判断", "腿1竞价占持仓‰", "标签", "预估页"],
            [[r["开盘"], r["场次"], r["组合"], move_cell(r), fmt(r["昨收价差"], 2),
              colored(r["针变化"], 2), r["针方向"], colored(r["首笔变化"], 2), r["首笔方向"], colored(r["腿1跳空"], 2),
              colored(r["主力跳空"], 2), r["针真假"], judged(r), fmt(r["腿1竞价占持仓‰"], 2), tags(r),
              Raw(f"<a href='{forecast_link(r['开盘'])}'>看</a>") if forecast_link(r["开盘"]) else ""]
             for _, r in recent.iterrows()]))
    PAGE_DIR.mkdir(exist_ok=True)
    f = PAGE_DIR / f"{rec['开盘'].max().replace(' ', '_').replace(':', '')}.html"
    f.write_text(page("开盘复盘", "".join(body)), encoding="utf-8")
    return f


def write_xlsx(df: pd.DataFrame, f: Path) -> None:
    """给人用 Excel 看的版本：列宽按内容自动调（中文算两格），开盘时间存成文字，不会出现 ####。
    程序读的还是同名 .csv。Excel 正开着这个文件时写不进去，只打印提示。"""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    def cell(v):
        if v is None or (isinstance(v, float) and math.isnan(v)):
            return None
        if isinstance(v, (float, np.floating)):
            v = float(v)
            return int(v) if v.is_integer() else v
        return v.item() if isinstance(v, np.generic) else v

    def width(v) -> int:
        return sum(2 if ord(ch) > 0x2E80 else 1 for ch in ("" if v is None else str(v)))

    wb = Workbook()
    ws = wb.active
    ws.title = f.stem
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([cell(v) for v in row])
    for i, col in enumerate(df.columns, 1):
        w = max([width(col)] + [width(c.value) for c in ws[get_column_letter(i)][1:]])
        ws.column_dimensions[get_column_letter(i)].width = min(w + 2, 80)
    ws.freeze_panes = "A2"
    try:
        wb.save(f)
    except PermissionError:
        print(f"{f.name} 正在 Excel 里开着，这次没更新（关掉后重跑 开盘复盘.py 就行）")


def run(out: Path = DATA) -> None:
    """重算所有开盘：写 复盘数据/开盘复盘.csv、更新 复盘数据/系数历史.csv 的实盘部分、写 复盘网页/开盘日期_时间.html。出错只打印，不抛出。"""
    if Path(out).resolve() != DATA.resolve():
        print("自定义取数目录：跳过正式开盘复盘，正式复盘记录、系数历史和网页不更新")
        return
    try:
        rec = build_records(out)
        REVIEW.mkdir(exist_ok=True)
        rec[OUT_COLS].to_csv(REC_FILE, index=False, encoding="utf-8-sig")
        write_xlsx(rec[OUT_COLS], REC_FILE.with_suffix(".xlsx"))
        live = coef_live(rec)
        update_coef(live)
        f = write_page(rec)
        print(f"开盘复盘：{rec['开盘'].nunique()} 次开盘、{len(rec)} 行；系数历史新增实盘 {len(live)} 行；网页 {f.name if f else '没写（还没有开过盘的数据）'}")
    except Exception as e:
        import traceback
        print(f"开盘复盘出错：{type(e).__name__} {e}\n{traceback.format_exc()}")


if __name__ == "__main__":
    run()
